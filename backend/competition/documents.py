"""生成、渲染并验证 3S 申报文档。"""

from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import uuid
import zipfile
from collections.abc import Callable
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

from backend.competition.document_support import (
    DocumentGenerationError,
    load_verified_evidence as _load_verified_evidence,
    publish_directory_atomically as _publish_directory,
    resolve_latest_evidence as _resolve_latest_evidence,
    sha256_file as _sha256,
)
from backend.competition.legacy_doc import (
    _pid_exists,
    _read_pid,
    _terminate_pid,
    _terminate_process,
)
from backend.competition.markdown_documents import DOCUMENT_IDS
from backend.competition.submission import (
    PROJECT_NAME,
    SubmissionGateError,
    load_verified_submission_metadata,
)


REPOSITORY_ROOT = Path(__file__).resolve().parents[2]
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
CP_NS = "http://schemas.openxmlformats.org/package/2006/metadata/core-properties"
DC_NS = "http://purl.org/dc/elements/1.1/"
NS = {"w": W_NS, "cp": CP_NS, "dc": DC_NS}

DESIGN_TOKENS: dict[str, object] = {
    "preset": "standard_business_brief",
    "headerTemplate": "memo_masthead",
    "page": {
        "size": "Letter portrait",
        "widthDxa": 12240,
        "heightDxa": 15840,
        "marginsDxa": 1440,
        "headerFooterDistanceDxa": 708,
        "usableWidthDxa": 9360,
    },
    "body": {
        "latinFont": "Calibri",
        "eastAsiaFontOverride": "Microsoft YaHei",
        "sizeHalfPoints": 22,
        "afterTwips": 120,
        "lineTwips": 264,
    },
    "headings": {
        "h1": {"sizeHalfPoints": 32, "beforeTwips": 320, "afterTwips": 160, "color": "2E74B5"},
        "h2": {"sizeHalfPoints": 26, "beforeTwips": 240, "afterTwips": 120, "color": "2E74B5"},
        "h3": {"sizeHalfPoints": 24, "beforeTwips": 160, "afterTwips": 80, "color": "1F4D78"},
    },
    "numbering": {
        "markerAlignedAtDxa": 360,
        "textIndentDxa": 720,
        "hangingDxa": 360,
        "afterTwips": 160,
        "lineTwips": 280,
    },
    "tables": {
        "widthDxa": 9360,
        "indentDxa": 120,
        "cellMarginsDxa": {"top": 80, "bottom": 80, "start": 120, "end": 120},
        "headerFill": "F2F4F7",
    },
}


WORD_PDF_DIAGNOSTIC_PATTERN = re.compile(
    r"(?m)^WORD_PDF_DIAGNOSTIC "
    r"stage=(?P<stage>[A-Za-z][A-Za-z0-9]*) "
    r"exceptionType=(?P<exception_type>[A-Za-z][A-Za-z0-9_.]*) "
    r"hresult=(?P<hresult>0x[0-9A-F]{8})$"
)


def _safe_word_pdf_diagnostic(stderr: str) -> str | None:
    match = WORD_PDF_DIAGNOSTIC_PATTERN.search(stderr)
    if match is None:
        return None
    return (
        f"stage={match.group('stage')}; "
        f"exceptionType={match.group('exception_type')}; "
        f"hresult={match.group('hresult')}"
    )


class WordPdfExporter:
    """通过可注入命令在隔离 Word 进程中导出 PDF。"""

    def __init__(
        self,
        *,
        command_builder: Callable[[Path, Path, Path], list[str]] | None = None,
        process_factory: Callable[..., subprocess.Popen[str]] = subprocess.Popen,
        terminate_pid: Callable[[int], None] = _terminate_pid,
        timeout_seconds: float = 60,
    ) -> None:
        self.command_builder = command_builder or _default_pdf_command
        self.process_factory = process_factory
        self.terminate_pid = terminate_pid
        self.timeout_seconds = timeout_seconds

    def export(self, source: Path, destination: Path) -> None:
        source = source.resolve()
        destination = destination.resolve()
        if not source.is_file() or source.suffix.lower() != ".docx":
            raise DocumentGenerationError("DOCX 转 PDF 的显式来源不存在。")
        if destination.exists() or destination.suffix.lower() != ".pdf":
            raise DocumentGenerationError("PDF 目标无效或已存在，拒绝覆盖。")
        destination.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(prefix=".word-pdf-", dir=destination.parent) as temporary_directory:
            pid_file = Path(temporary_directory) / "word.pid"
            process: subprocess.Popen[str] | None = None
            word_pid: int | None = None
            try:
                process = self.process_factory(
                    self.command_builder(source, destination, pid_file),
                    stdin=subprocess.DEVNULL,
                    stdout=subprocess.PIPE,
                    stderr=subprocess.PIPE,
                    text=True,
                    encoding="utf-8",
                    errors="replace",
                    creationflags=(subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0),
                )
                try:
                    _, stderr = process.communicate(timeout=self.timeout_seconds)
                except subprocess.TimeoutExpired as exc:
                    word_pid = _read_pid(pid_file)
                    _terminate_process(process)
                    if word_pid is not None and word_pid != process.pid:
                        self.terminate_pid(word_pid)
                    raise DocumentGenerationError("DOCX 转 PDF 超时，已清理隔离进程。") from exc
                word_pid = _read_pid(pid_file)
                if process.returncode != 0:
                    diagnostic = _safe_word_pdf_diagnostic(stderr)
                    if diagnostic is None:
                        raise DocumentGenerationError("隔离 Word 未能导出 PDF。")
                    raise DocumentGenerationError(
                        f"隔离 Word 未能导出 PDF（{diagnostic}）。"
                    )
                if not destination.is_file() or not destination.read_bytes().startswith(b"%PDF"):
                    raise DocumentGenerationError("隔离 Word 未生成有效 PDF。")
            except DocumentGenerationError:
                destination.unlink(missing_ok=True)
                raise
            except (OSError, subprocess.SubprocessError) as exc:
                destination.unlink(missing_ok=True)
                raise DocumentGenerationError("DOCX 转 PDF 进程无法安全完成。") from exc
            finally:
                if process is not None and process.poll() is None:
                    _terminate_process(process)
                if word_pid is None:
                    word_pid = _read_pid(pid_file)
                if word_pid is not None and (process is None or word_pid != process.pid) and _pid_exists(word_pid):
                    self.terminate_pid(word_pid)


def _default_pdf_command(source: Path, destination: Path, pid_file: Path) -> list[str]:
    executable = os.environ.get("WAREHOUSE_PATROL_PWSH") or shutil.which("pwsh")
    worker = REPOSITORY_ROOT / "competition/docx-to-pdf-worker.ps1"
    if executable is None:
        raise DocumentGenerationError("找不到 PowerShell 7，无法启动隔离 Word。")
    if not worker.is_file():
        raise DocumentGenerationError("DOCX 转 PDF 隔离脚本不存在。")
    return [
        executable, "-NoLogo", "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass",
        "-File", str(worker), "-SourcePath", str(source), "-DestinationPath", str(destination),
        "-PidFile", str(pid_file),
    ]


def _external_command(executable: Path, arguments: list[str]) -> list[str]:
    if executable.suffix.lower() in {".cmd", ".bat"}:
        return ["cmd.exe", "/d", "/s", "/c", str(executable), *arguments]
    return [str(executable), *arguments]


def _render_pdf_pages(
    document_id: str,
    pdf_path: Path,
    output_directory: Path,
    pdftoppm: Path,
    pdfinfo: Path,
) -> list[dict[str, object]]:
    if not pdftoppm.is_file() or not pdfinfo.is_file():
        raise DocumentGenerationError("Poppler 的 pdftoppm/pdfinfo 路径无效。")
    info = subprocess.run(
        _external_command(pdfinfo, [str(pdf_path)]),
        check=False,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    if info.returncode != 0:
        raise DocumentGenerationError("pdfinfo 无法读取导出的 PDF。")
    match = re.search(r"(?mi)^Pages:\s*(\d+)\s*$", info.stdout)
    if match is None or int(match.group(1)) < 1:
        raise DocumentGenerationError("pdfinfo 未返回有效页数。")
    page_count = int(match.group(1))
    page_directory = output_directory / "rendered" / document_id
    page_directory.mkdir(parents=True, exist_ok=True)
    prefix = page_directory / "page"
    rendered = subprocess.run(
        _external_command(pdftoppm, ["-png", "-r", "144", str(pdf_path), str(prefix)]),
        check=False,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    if rendered.returncode != 0:
        raise DocumentGenerationError("pdftoppm 无法渲染导出的 PDF。")
    pages: list[dict[str, object]] = []
    expected_names: set[str] = set()
    for page_number in range(1, page_count + 1):
        source = page_directory / f"page-{page_number}.png"
        target = page_directory / f"page-{page_number:03d}.png"
        if not source.is_file():
            raise DocumentGenerationError("pdftoppm 未生成全部 PDF 页图。")
        source.replace(target)
        expected_names.add(target.name)
        pages.append(
            {
                "documentId": document_id,
                "pageNumber": page_number,
                "imagePath": target.relative_to(output_directory).as_posix(),
                "imageSha256": _sha256(target),
                "status": "pending",
            }
        )
    if {item.name for item in page_directory.glob("page-*.png")} != expected_names:
        raise DocumentGenerationError("pdftoppm 页图数量与 pdfinfo 不一致。")
    return pages


def verify_visual_approvals(output_directory: Path, approvals_path: Path | None) -> bool:
    if approvals_path is None or not approvals_path.is_file():
        return False
    request_path = output_directory / "visual-review-request.json"
    if not request_path.is_file():
        return False
    try:
        request = json.loads(request_path.read_text(encoding="utf-8"))
        approvals = json.loads(approvals_path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError):
        return False
    if not isinstance(approvals, dict) or approvals.get("schemaVersion") != 1 or set(approvals) != {"schemaVersion", "pages"}:
        return False
    requested_pages = [
        page
        for document in request.get("documents", {}).values()
        if isinstance(document, dict)
        for page in document.get("pages", [])
        if isinstance(page, dict)
    ]
    approval_pages = approvals.get("pages")
    if not requested_pages or not isinstance(approval_pages, list):
        return False
    expected: dict[tuple[str, int], str] = {}
    for page in requested_pages:
        key = (page.get("documentId"), page.get("pageNumber"))
        image_path = page.get("imagePath")
        digest = page.get("imageSha256")
        if not isinstance(key[0], str) or not isinstance(key[1], int) or not isinstance(image_path, str) or not isinstance(digest, str) or key in expected:
            return False
        image = output_directory / image_path
        if not image.is_file() or _sha256(image) != digest:
            return False
        expected[key] = digest
    actual: dict[tuple[str, int], str] = {}
    for approval in approval_pages:
        if not isinstance(approval, dict) or set(approval) != {"documentId", "pageNumber", "imageSha256", "approved"}:
            return False
        key = (approval.get("documentId"), approval.get("pageNumber"))
        digest = approval.get("imageSha256")
        if key in actual or approval.get("approved") is not True or not isinstance(key[0], str) or not isinstance(key[1], int) or not isinstance(digest, str):
            return False
        actual[key] = digest
    return actual == expected


def _set_run_font(
    run: Any,
    *,
    size: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    run.font.name = "Calibri"
    r_fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), "Calibri")
    r_fonts.set(qn("w:hAnsi"), "Calibri")
    r_fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def _configure_styles(document: Any) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    def style_font(style: Any, size: float, color: str, bold: bool) -> None:
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = bold
        r_pr = style.element.get_or_add_rPr()
        r_fonts = r_pr.find(qn("w:rFonts"))
        if r_fonts is None:
            r_fonts = OxmlElement("w:rFonts")
            r_pr.insert(0, r_fonts)
        r_fonts.set(qn("w:ascii"), "Calibri")
        r_fonts.set(qn("w:hAnsi"), "Calibri")
        r_fonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    normal = document.styles["Normal"]
    style_font(normal, 11, "000000", False)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    title = document.styles["Title"]
    style_font(title, 23, "000000", True)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(4)
    title.paragraph_format.line_spacing = 1.0

    subtitle = document.styles["Subtitle"]
    style_font(subtitle, 14, "373737", False)
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(16)
    subtitle.paragraph_format.line_spacing = 1.0

    for name, size, color, before, after in (
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ):
        style = document.styles[name]
        style_font(style, size, color, True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.keep_with_next = True


def _configure_section(document: Any) -> None:
    from docx.shared import Inches

    for section in document.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.right_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)


def _add_numbering(document: Any) -> int:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    numbering = document.part.numbering_part.element
    abstract_ids = [
        int(item.get(qn("w:abstractNumId")))
        for item in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [int(item.get(qn("w:numId"))) for item in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    for tag, value in (("w:start", "1"), ("w:numFmt", "decimal"), ("w:lvlText", "%1."), ("w:lvlJc", "left")):
        node = OxmlElement(tag)
        node.set(qn("w:val"), value)
        level.append(node)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    p_pr.append(tabs)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "160")
    spacing.set(qn("w:line"), "280")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "720")
    indent.set(qn("w:hanging"), "360")
    p_pr.append(indent)
    level.append(p_pr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def _numbered_paragraph(document: Any, text: str, num_id: int) -> Any:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    paragraph = document.add_paragraph(style="Normal")
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    identifier = OxmlElement("w:numId")
    identifier.set(qn("w:val"), str(num_id))
    num_pr.extend((ilvl, identifier))
    paragraph._p.get_or_add_pPr().append(num_pr)
    _set_run_font(paragraph.add_run(text))
    return paragraph


def _set_table_geometry(table: Any, widths: list[int]) -> None:
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    if sum(widths) != 9360:
        raise DocumentGenerationError("表格列宽必须精确合计 9360 DXA。")
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    for tag in ("w:tblW", "w:tblInd", "w:tblLayout", "w:tblCellMar", "w:tblBorders"):
        existing = tbl_pr.find(qn(tag))
        if existing is not None:
            tbl_pr.remove(existing)
    tbl_width = OxmlElement("w:tblW")
    tbl_width.set(qn("w:w"), "9360")
    tbl_width.set(qn("w:type"), "dxa")
    tbl_pr.insert(0, tbl_width)
    indent = OxmlElement("w:tblInd")
    indent.set(qn("w:w"), "120")
    indent.set(qn("w:type"), "dxa")
    tbl_pr.append(indent)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    margins = OxmlElement("w:tblCellMar")
    for side, value in (("top", 80), ("start", 120), ("bottom", 80), ("end", 120)):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    tbl_pr.append(margins)
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "start", "bottom", "end", "insideH", "insideV"):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), "D0D7DE")
        borders.append(node)
    tbl_pr.append(borders)

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_width = tc_pr.find(qn("w:tcW"))
            if tc_width is None:
                tc_width = OxmlElement("w:tcW")
                tc_pr.insert(0, tc_width)
            tc_width.set(qn("w:w"), str(width))
            tc_width.set(qn("w:type"), "dxa")


def _add_table(document: Any, headers: list[str], rows: list[list[str]], widths: list[int]) -> Any:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    table = document.add_table(rows=1, cols=len(headers))
    for index, text in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.style = document.styles["Normal"]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_run_font(paragraph.add_run(text), bold=True)
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "F2F4F7")
        cell._tc.get_or_add_tcPr().append(shading)
    for values in rows:
        cells = table.add_row().cells
        for index, text in enumerate(values):
            cells[index].text = ""
            paragraph = cells[index].paragraphs[0]
            paragraph.style = document.styles["Normal"]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            _set_run_font(paragraph.add_run(text))
    _set_table_geometry(table, widths)
    document.add_paragraph(style="Normal")
    return table


def _set_header_footer(document: Any, label: str) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    section = document.sections[0]
    header = section.header.paragraphs[0]
    header.text = ""
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header.paragraph_format.space_after = 0
    _set_run_font(header.add_run(f"仓巡智调 | {label}"), size=9, color="666666")
    footer = section.footer.paragraphs[0]
    footer.text = ""
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_before = 0
    _set_run_font(footer.add_run("第 "), size=9, color="666666")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    _set_run_font(footer.add_run(" 页"), size=9, color="666666")


def _masthead(document: Any, title: str, label: str, metadata_rows: list[tuple[str, str]]) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    _set_header_footer(document, label)
    title_paragraph = document.add_paragraph(style="Title")
    _set_run_font(title_paragraph.add_run(title), size=23, bold=True)
    subtitle = document.add_paragraph(style="Subtitle")
    _set_run_font(subtitle.add_run(PROJECT_NAME), size=14, color="373737")
    for key, value in metadata_rows:
        paragraph = document.add_paragraph(style="Normal")
        paragraph.paragraph_format.space_after = Pt(2)
        _set_run_font(paragraph.add_run(f"{key}："), bold=True)
        _set_run_font(paragraph.add_run(value))
    rule = document.add_paragraph()
    rule.paragraph_format.space_before = Pt(6)
    rule.paragraph_format.space_after = Pt(8)
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:color"), "2E74B5")
    p_bdr.append(bottom)
    rule._p.get_or_add_pPr().append(p_bdr)


def _paragraph(document: Any, text: str, *, bold_label: str | None = None) -> Any:
    paragraph = document.add_paragraph(style="Normal")
    if bold_label is not None:
        _set_run_font(paragraph.add_run(bold_label), bold=True)
    _set_run_font(paragraph.add_run(text))
    return paragraph


def _new_document(document_id: str, title: str, metadata_rows: list[tuple[str, str]]) -> tuple[Any, int]:
    from docx import Document

    document = Document()
    _configure_section(document)
    _configure_styles(document)
    document.core_properties.creator = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.title = title
    document.core_properties.subject = document_id
    document.core_properties.keywords = "3S, warehouse patrol, simulation"
    _masthead(document, title, document_id, metadata_rows)
    return document, _add_numbering(document)


def _build_application(metadata: dict[str, object]) -> Any:
    applicant = metadata["applicant"]
    advisors = metadata["advisors"]
    procedures = metadata["procedures"]
    materials = metadata["requiredMaterials"]
    document, _ = _new_document(
        "application-summary",
        "申报信息汇总",
        [
            ("材料性质", "非官方申报表；仅汇总已核实本地元数据"),
            ("申请人", applicant["name"]),
            ("学校与专业", f"{applicant['school']} / {applicant['major']}"),
            ("核验状态", "外部手续元数据严格门已通过"),
        ],
    )
    document.add_heading("使用边界", level=1)
    _paragraph(
        document,
        "当前源码工作区没有官方 .doc、.docx 或 .pdf 申报模板。本文件是非官方申报表，"
        "不得冒充主办方表格；若正式规则要求官方表格，必须由用户提供被 Git 忽略的显式本地原件，"
        "再通过只读转换流程处理。",
    )
    document.add_heading("申请人信息", level=1)
    _add_table(
        document,
        ["字段", "已核实内容"],
        [
            ["姓名", applicant["name"]],
            ["学号", applicant["studentId"]],
            ["联系方式", f"{applicant['phone']} / {applicant['email']}"],
            ["学校", applicant["school"]],
            ["专业", applicant["major"]],
        ],
        [2700, 6660],
    )
    document.add_heading("指导教师或推荐者", level=1)
    _add_table(
        document,
        ["姓名", "角色", "职称"],
        [[item["name"], item["role"], item["professionalTitle"]] for item in advisors],
        [3000, 3000, 3360],
    )
    document.add_heading("外部手续核验", level=1)
    labels = {
        "eligibility": "资格",
        "majorAndSchoolAcceptance": "专业与学校受理",
        "advisorAndRecommenderRequirements": "指导教师/推荐者角色与职称",
        "collegeAndSchoolRecommendation": "学院/学校推荐",
        "signaturesAndSeals": "签字盖章",
        "internalDeadline": "校内截止",
        "registrationSystem": "报名系统",
        "materialsMailbox": "材料邮箱",
        "attachmentConstraints": "附件大小与格式",
    }
    rows = []
    for key, label in labels.items():
        item = procedures[key]
        value = item["statement"]
        if key == "internalDeadline":
            value += f" 截止时间：{item['deadline']}。"
        elif key == "registrationSystem":
            value += f" 系统：{item['systemUrl']}。"
        elif key == "materialsMailbox":
            value += f" 邮箱：{item['email']}。"
        elif key == "attachmentConstraints":
            value += f" 最大 {item['maximumBytes']} 字节；格式：{', '.join(item['allowedFormats'])}。"
        rows.append([label, value, item["source"]["url"]])
    _add_table(document, ["核验项", "书面要求", "来源 URL"], rows, [1600, 4760, 3000])
    document.add_heading("必交材料关系", level=1)
    material_labels = {
        "sourceCode": "源码",
        "runnablePackage": "运行包",
        "video": "视频",
        "applicationForm": "申报书",
        "technicalReport": "技术报告",
    }
    _add_table(
        document,
        ["材料", "是否必交", "已核实关系"],
        [
            [material_labels[key], "是" if item["required"] else "否", item["relationship"]]
            for key, item in materials.items()
        ],
        [1500, 1500, 6360],
    )
    return document


def _evidence_run(evidence: dict[str, object], case_id: str) -> dict[str, object] | None:
    return next((run for run in evidence["runs"] if run.get("caseId") == case_id), None)


def _build_technical_report(evidence: dict[str, object]) -> Any:
    document, num_id = _new_document(
        "technical-report",
        "技术报告",
        [
            ("系统形态", "本地 Windows x64 软件仿真系统"),
            ("证据口径", "可重算固定案例与在线流程证据"),
            ("边界声明", "不声称完整 CBS/MAPF 或任意输入零冲突"),
        ],
    )
    sections = [
        (
            "问题背景",
            "仓储巡检与搬运任务具有持续到达、优先级变化、临时阻塞和机器人故障等动态特征。"
            "系统目标是在网格化软件仿真中统一处理巡检、取送与紧急任务，并保留可解释的计划、事件和指标。",
        ),
        (
            "系统架构",
            "前端采用 React、TypeScript 与 Vite 展示地图、路径、队列、事件和指标；后端采用 FastAPI 维护调度接口与在线会话。"
            "当前在线状态是进程内会话，只保证单进程并发边界；系统无持久化、无认证，也未接入实体物联网、低空通信或云平台。",
        ),
        (
            "任务分配",
            "后端依据任务类型能力、载荷、电量、可达性、优先级、锁定前缀和滚动窗口进行候选分配。"
            "该方法属于确定性启发式与 beam-style 组合搜索，不使用机器学习。",
        ),
        (
            "时空规划",
            "单机器人静态路径使用 A*；多机器人执行通过按优先级的时间占用与预留约束降低顶点和边冲突，"
            "并在候选计划中检查冲突、期限和失败。该实现不是完整 CBS，也不提供完整 MAPF 保证。",
        ),
        (
            "安全执行",
            "在线移动以最终 timed path 为权威；提交每个 tick 前执行全车安全门和能量门。"
            "重复安全干预会形成可观察的 safety stall，但不会通过削弱全车安全等待来掩盖冲突。",
        ),
        (
            "动态恢复",
            "系统在任务到达、阻塞单元、机器人故障与恢复、滚动窗口纳入和场景动态触发时重规划。"
            "失败详情区分永久不可达与可通过清障或恢复机器人解决的临时条件，并保持部分任务进度。",
        ),
    ]
    for heading, body in sections:
        document.add_heading(heading, level=1)
        _paragraph(document, body)

    document.add_heading("实验方法与结果", level=1)
    _paragraph(
        document,
        "实验使用同一源码中的固定场景、固定种子压力案例、主演示在线流程和安全门边界案例。"
        "每条结果均要求 outcome=completed 且 accepted=true；原始 JSON/CSV、图表和证据清单应随发布包保留，"
        "以便独立重算，而不是把页面中的摘要当作唯一证据。",
    )
    without = _evidence_run(evidence, "integrated-demo-without-conflict-avoidance")
    with_avoidance = _evidence_run(evidence, "integrated-demo-with-conflict-avoidance")
    main = _evidence_run(evidence, "main-demo-online")
    safety = _evidence_run(evidence, "safety-gate-boundary")
    result_rows: list[list[str]] = []
    if without and with_avoidance:
        result_rows.append(
            [
                "避碰对照",
                f"预测冲突 {without['metrics']['conflictCount']} -> {with_avoidance['metrics']['conflictCount']}",
                "同一 integrated-demo；不外推到任意输入",
            ]
        )
    if main:
        result_rows.append(
            ["在线流程", f"完成任务 {main['metrics']['completedTaskCount']}", "固定主演示流程"],
        )
    if safety:
        result_rows.append(
            [
                "安全门",
                f"干预 {safety['metrics']['safetyInterventionCount']} 次，stall={safety['metrics']['safetyStallCount']}",
                "受控边界案例",
            ]
        )
    pressure_runs = [run for run in evidence["runs"] if str(run.get("caseId", "")).startswith("seeded-pressure")]
    for run in pressure_runs:
        result_rows.append(
            [
                "固定种子压力",
                f"{run.get('robotCount')} 台机器人 / {run.get('taskCount')} 项任务 / {run['metrics'].get('replanTimeMs')} ms",
                str(run.get("caseId")),
            ]
        )
    _add_table(document, ["案例", "受控结果", "适用范围"], result_rows, [1900, 3860, 3600])

    document.add_heading("同类技术比较", level=1)
    _add_table(
        document,
        ["技术路线", "优势", "与本系统的关系"],
        [
            ["静态一次性调度", "实现简单、批次边界清晰", "本系统增加在线任务流、状态保持与恢复"],
            ["优先级时间预留", "工程可解释、适合滚动规划", "本系统采用此类思路并加入安全执行门"],
            ["完整 CBS/MAPF 求解", "在严格模型下可提供更强完备性或最优性", "本系统未实现，不作等价宣称"],
            ["学习型调度", "可从数据中拟合策略", "本系统不使用机器学习，避免无证据泛化"],
        ],
        [2100, 3300, 3960],
    )
    document.add_heading("应用价值", level=1)
    _paragraph(
        document,
        "系统可作为仓储巡检、校园或设施巡逻等场景的调度算法验证与操作演示平台，"
        "用于观察任务分配、路径、安全干预、动态恢复和指标变化。其价值首先体现在可运行、可解释和可重算的软件仿真。",
    )
    document.add_heading("限制", level=1)
    for text in (
        "仅完成软件仿真，未完成实体机器人、传感器、低空通信、云平台或真实仓库部署。",
        "在线状态为进程内会话，存在单进程并发边界，无持久化、无认证。",
        "采用优先级与预留式规划，不提供完整 MAPF 保证，也不声称任意输入零冲突。",
        "性能阈值与实验耗时只代表已记录机器和固定案例，不代表跨机器普适上界。",
    ):
        _numbered_paragraph(document, text, num_id)
    document.add_heading("参考文献", level=1)
    for reference in (
        "Hart, P. E., Nilsson, N. J., and Raphael, B. A Formal Basis for the Heuristic Determination of Minimum Cost Paths. IEEE Transactions on Systems Science and Cybernetics, 1968.",
        "Silver, D. Cooperative Pathfinding. Proceedings of AIIDE, 2005.",
        "Stern, R. et al. Multi-Agent Pathfinding: Definitions, Variants, and Benchmarks. SoCS, 2019.",
    ):
        _numbered_paragraph(document, reference, num_id)
    document.add_heading("知识产权", level=1)
    _paragraph(
        document,
        "项目源码、文档与实验脚本的自研部分按仓库权属说明管理；第三方运行库、构建工具与前端依赖不转化为自研成果，"
        "其名称、锁定版本和锁文件可见许可信息另列于《知识产权与第三方依赖清单》。未在锁文件声明的许可不得猜填。",
    )
    return document


def _build_run_guide() -> Any:
    document, num_id = _new_document(
        "run-guide",
        "运行说明",
        [
            ("交付形态", "Windows x64 onedir 便携目录"),
            ("默认监听", "127.0.0.1 本机回环地址"),
            ("目标机依赖", "便携包运行时不要求另装 Node.js 或 Python"),
        ],
    )
    document.add_heading("启动步骤", level=1)
    for text in (
        "完整解压 Windows x64 便携目录，保留 WarehousePatrol.exe 与 _internal 的相对结构。",
        "双击 WarehousePatrol.exe；自动化验收可使用 WarehousePatrol.exe --headless --port 8123。",
        "浏览器访问启动器给出的本机地址，确认后端状态为在线。",
        "创建在线会话，按需添加任务、触发动态事件、推进 tick，并查看路径、事件和指标。",
        "结束后关闭桌面窗口；无头模式使用 Ctrl+Break，使进程正常清理并释放端口。",
    ):
        _numbered_paragraph(document, text, num_id)
    document.add_heading("主要操作", level=1)
    _add_table(
        document,
        ["操作", "预期结果", "注意事项"],
        [
            ["创建会话", "初始化机器人、任务、库存和时间线", "历史回放时不允许运行态修改"],
            ["推进 tick", "机器人按 timed path 执行", "安全门或能量门可能保持等待"],
            ["添加任务", "统一任务端点接收手工或生成任务", "后端校验能力、载荷与输入上限"],
            ["阻塞/故障", "触发重规划并暴露恢复建议", "恢复动作只处理实际可恢复条件"],
        ],
        [1800, 3660, 3900],
    )
    document.add_heading("运行边界", level=1)
    _paragraph(
        document,
        "本交付是软件仿真。会话保存在单个服务进程内，只有单进程并发保证；服务重启后无持久化恢复，且无认证。"
        "规划不提供完整 MAPF 保证。未知占用端口不会被启动器终止或接管。",
    )
    document.add_heading("故障排查", level=1)
    _add_table(
        document,
        ["现象", "检查", "处理"],
        [
            ["端口已占用", "查看启动器给出的端口错误", "关闭已知占用程序或改用新的显式端口"],
            ["页面无法打开", "确认进程仍在且 /health 返回 ok", "重新启动便携包；不要删除 _internal"],
            ["任务未分配", "查看 failureCategory、blockingCells、blockingRobotIds", "按 recoveryAction 清障或恢复机器人"],
            ["播放暂停", "查看 tick 的业务错误详情", "修正输入后继续，不把 4xx 误判为后端离线"],
        ],
        [1900, 3860, 3600],
    )
    return document


def _dependency_rows() -> list[list[str]]:
    frontend_package = json.loads((REPOSITORY_ROOT / "frontend/package.json").read_text(encoding="utf-8"))
    frontend_lock = json.loads((REPOSITORY_ROOT / "frontend/package-lock.json").read_text(encoding="utf-8"))
    rows: list[list[str]] = []
    for group in ("dependencies", "devDependencies"):
        for name, requested in sorted(frontend_package.get(group, {}).items()):
            locked = frontend_lock.get("packages", {}).get(f"node_modules/{name}", {})
            rows.append(["npm", name, str(locked.get("version", requested)), str(locked.get("license", "锁文件未声明"))])
    for line in (REPOSITORY_ROOT / "backend/requirements.lock.txt").read_text(encoding="utf-8").splitlines():
        stripped = line.strip()
        if not stripped or stripped.startswith("#"):
            continue
        name, _, version = stripped.partition("==")
        rows.append(["PyPI", name, version or "锁文件未声明版本", "锁文件未声明"])
    return rows


def _build_dependency_inventory() -> Any:
    document, _ = _new_document(
        "third-party-dependencies",
        "知识产权与第三方依赖清单",
        [
            ("清单来源", "仓库锁文件与 package manifest"),
            ("许可口径", "仅记录锁文件显式值；缺失时不猜填"),
            ("素材边界", "正式发布仅组装已授权源码、软件、证据、文档与视频"),
        ],
    )
    document.add_heading("自研与第三方边界", level=1)
    _paragraph(
        document,
        "仓巡智调的业务源码、测试、竞赛清单和材料生成代码属于本项目交付范围。"
        "Python、FastAPI、Uvicorn、React、Vite、PyInstaller、python-docx、Poppler 和其他依赖仍受各自许可证约束，"
        "不会因打包进入便携目录而被描述为自研成果。",
    )
    document.add_heading("依赖清单", level=1)
    _add_table(
        document,
        ["生态", "依赖", "锁定版本", "锁文件许可字段"],
        _dependency_rows(),
        [1000, 3000, 1800, 3560],
    )
    document.add_heading("许可与素材核对要求", level=1)
    _paragraph(
        document,
        "表中“锁文件未声明”是对输入文件的客观描述，不是许可结论。正式提交前应保留上游许可证文本或书面授权；"
        "敏感信息扫描会拒绝本地元数据、.env、缓存、个人绝对路径和未进入授权组装清单的素材。",
    )
    return document


def _scrub_docx_metadata(path: Path) -> None:
    temporary = path.with_name(f".{path.name}.{uuid.uuid4().hex}.tmp")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temporary, "w") as target:
        for info in source.infolist():
            if info.filename == "docProps/custom.xml":
                continue
            data = source.read(info.filename)
            if info.filename == "docProps/core.xml":
                root = ElementTree.fromstring(data)
                for tag in (f"{{{DC_NS}}}creator", f"{{{CP_NS}}}lastModifiedBy"):
                    node = root.find(tag)
                    if node is not None:
                        node.text = ""
                data = ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)
            elif info.filename.startswith("word/") and info.filename.endswith(".xml"):
                try:
                    from lxml import etree
                except ModuleNotFoundError as exc:
                    raise DocumentGenerationError(
                        "文档 Python 缺少锁定的 lxml 依赖。"
                    ) from exc
                parser = etree.XMLParser(
                    resolve_entities=False,
                    no_network=True,
                    load_dtd=False,
                    remove_blank_text=False,
                    strip_cdata=False,
                )
                root = etree.fromstring(data, parser=parser)
                for node in root.iter():
                    for attribute in list(node.attrib):
                        if attribute.startswith(f"{{{W_NS}}}rsid"):
                            del node.attrib[attribute]
                document_tree = root.getroottree()
                serialize_options: dict[str, object] = {
                    "encoding": "UTF-8",
                    "xml_declaration": True,
                }
                if document_tree.docinfo.standalone:
                    serialize_options["standalone"] = True
                data = etree.tostring(document_tree, **serialize_options)
            target.writestr(info, data)
    os.replace(temporary, path)


def _attribute(node: ElementTree.Element | None, name: str) -> str | None:
    return None if node is None else node.get(f"{{{W_NS}}}{name}")


def audit_standard_business_brief(path: Path) -> list[str]:
    issues: list[str] = []
    try:
        with zipfile.ZipFile(path) as archive:
            styles = ElementTree.fromstring(archive.read("word/styles.xml"))
            document = ElementTree.fromstring(archive.read("word/document.xml"))
            numbering = ElementTree.fromstring(archive.read("word/numbering.xml"))
            core = ElementTree.fromstring(archive.read("docProps/core.xml"))
            names = set(archive.namelist())
    except (OSError, KeyError, zipfile.BadZipFile, ElementTree.ParseError):
        return ["DOCX 结构不可读"]

    if "docProps/custom.xml" in names:
        issues.append("仍包含自定义属性")
    for tag in (f"{{{DC_NS}}}creator", f"{{{CP_NS}}}lastModifiedBy"):
        node = core.find(tag)
        if node is not None and (node.text or "").strip():
            issues.append("核心作者元数据未清除")

    expected_styles = {
        "Normal": ("22", "000000", "0", "120", "264"),
        "Title": ("46", "000000", "0", "80", "240"),
        "Subtitle": ("28", "373737", "0", "320", "240"),
        "Heading1": ("32", "2E74B5", "320", "160", "240"),
        "Heading2": ("26", "2E74B5", "240", "120", "240"),
        "Heading3": ("24", "1F4D78", "160", "80", "240"),
    }
    for style_id, (size, color, before, after, line) in expected_styles.items():
        style = styles.find(f"w:style[@w:styleId='{style_id}']", NS)
        if style is None:
            issues.append(f"缺少样式 {style_id}")
            continue
        r_fonts = style.find("w:rPr/w:rFonts", NS)
        font_size = style.find("w:rPr/w:sz", NS)
        font_color = style.find("w:rPr/w:color", NS)
        spacing = style.find("w:pPr/w:spacing", NS)
        if (
            _attribute(r_fonts, "ascii") != "Calibri"
            or _attribute(r_fonts, "hAnsi") != "Calibri"
            or _attribute(r_fonts, "eastAsia") != "Microsoft YaHei"
            or _attribute(font_size, "val") != size
            or _attribute(font_color, "val") != color
            or _attribute(spacing, "before") != before
            or _attribute(spacing, "after") != after
            or _attribute(spacing, "line") != line
        ):
            issues.append(f"样式 {style_id} 未匹配精确令牌")

    section = document.find(".//w:sectPr", NS)
    size = None if section is None else section.find("w:pgSz", NS)
    margin = None if section is None else section.find("w:pgMar", NS)
    if _attribute(size, "w") != "12240" or _attribute(size, "h") != "15840":
        issues.append("页面尺寸不是 Letter portrait")
    if margin is None or any(_attribute(margin, side) != "1440" for side in ("top", "right", "bottom", "left")):
        issues.append("页边距不是 1440 DXA")
    if _attribute(margin, "header") != "708" or _attribute(margin, "footer") != "708":
        issues.append("页眉页脚距离不是 708 DXA")

    valid_numbering = False
    for level in numbering.findall("w:abstractNum/w:lvl", NS):
        fmt = level.find("w:numFmt", NS)
        indent = level.find("w:pPr/w:ind", NS)
        spacing = level.find("w:pPr/w:spacing", NS)
        if (
            _attribute(fmt, "val") == "decimal"
            and _attribute(indent, "left") == "720"
            and _attribute(indent, "hanging") == "360"
            and _attribute(spacing, "after") == "160"
            and _attribute(spacing, "line") == "280"
        ):
            valid_numbering = True
            break
    if not valid_numbering:
        issues.append("缺少精确实编号定义")

    for table_index, table in enumerate(document.findall(".//w:tbl", NS), start=1):
        width = table.find("w:tblPr/w:tblW", NS)
        indent = table.find("w:tblPr/w:tblInd", NS)
        grid_widths = [int(_attribute(item, "w") or "-1") for item in table.findall("w:tblGrid/w:gridCol", NS)]
        if _attribute(width, "w") != "9360" or _attribute(width, "type") != "dxa":
            issues.append(f"表格 {table_index} 总宽度不正确")
        if _attribute(indent, "w") != "120" or sum(grid_widths) != 9360:
            issues.append(f"表格 {table_index} 缩进或网格不正确")
        for row in table.findall("w:tr", NS):
            cell_widths = [int(_attribute(cell.find("w:tcPr/w:tcW", NS), "w") or "-1") for cell in row.findall("w:tc", NS)]
            if cell_widths != grid_widths:
                issues.append(f"表格 {table_index} 单元格宽度与网格不一致")
                break
    for paragraph in document.findall(".//w:p", NS):
        text = "".join(node.text or "" for node in paragraph.findall(".//w:t", NS)).lstrip()
        if text.startswith(("•", "●", "- ")):
            issues.append("发现假项目符号")
            break
    return issues


def generate_documents(
    *,
    metadata_path: Path,
    evidence_directory: Path,
    output_directory: Path,
    pdftoppm: Path,
    pdfinfo: Path,
    approvals_path: Path | None = None,
    word_timeout_seconds: float = 60,
) -> dict[str, object]:
    metadata = load_verified_submission_metadata(metadata_path, REPOSITORY_ROOT)
    evidence_path = _resolve_latest_evidence(evidence_directory)
    evidence = _load_verified_evidence(evidence_path)
    staging = output_directory.with_name(f".{output_directory.name}.{uuid.uuid4().hex}.tmp")
    if staging.exists():
        shutil.rmtree(staging)
    staging.mkdir(parents=True)
    exporter = WordPdfExporter(timeout_seconds=word_timeout_seconds)
    documents = {
        "application-summary": _build_application(metadata),
        "technical-report": _build_technical_report(evidence),
        "run-guide": _build_run_guide(),
        "third-party-dependencies": _build_dependency_inventory(),
    }
    request_documents: dict[str, object] = {}
    file_hashes: dict[str, str] = {}
    try:
        for document_id, document in documents.items():
            docx_path = staging / f"{document_id}.docx"
            pdf_path = staging / f"{document_id}.pdf"
            document.save(docx_path)
            _scrub_docx_metadata(docx_path)
            issues = audit_standard_business_brief(docx_path)
            if issues:
                raise DocumentGenerationError("DOCX 设计令牌审计失败：" + "；".join(issues))
            try:
                exporter.export(docx_path, pdf_path)
            except DocumentGenerationError as exc:
                raise DocumentGenerationError(
                    f"documentId={document_id}；{exc}"
                ) from exc
            pages = _render_pdf_pages(document_id, pdf_path, staging, pdftoppm, pdfinfo)
            request_documents[document_id] = {
                "pdfPath": pdf_path.relative_to(staging).as_posix(),
                "pages": pages,
            }
            file_hashes[docx_path.name] = _sha256(docx_path)
            file_hashes[pdf_path.name] = _sha256(pdf_path)
        request = {"schemaVersion": 1, "documents": request_documents}
        (staging / "visual-review-request.json").write_text(
            json.dumps(request, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        manifest = {
            "schemaVersion": 1,
            "projectName": PROJECT_NAME,
            "designTokens": DESIGN_TOKENS,
            "evidence": {"resultsSha256": _sha256(evidence_path / "results.json")},
            "files": dict(sorted(file_hashes.items())),
            "visualApproved": verify_visual_approvals(staging, approvals_path),
        }
        (staging / "document-manifest.json").write_text(
            json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        _publish_directory(staging, output_directory)
        return manifest
    except Exception:
        shutil.rmtree(staging, ignore_errors=True)
        raise


def _verify_existing_documents(output_directory: Path, approvals_path: Path | None) -> bool:
    for document_id in DOCUMENT_IDS:
        docx_path = output_directory / f"{document_id}.docx"
        pdf_path = output_directory / f"{document_id}.pdf"
        if not docx_path.is_file() or audit_standard_business_brief(docx_path):
            return False
        if not pdf_path.is_file() or not pdf_path.read_bytes().startswith(b"%PDF"):
            return False
    return verify_visual_approvals(output_directory, approvals_path)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="生成并验证 3S 申报材料")
    parser.add_argument("--metadata", required=True, type=Path)
    parser.add_argument(
        "--evidence-dir",
        type=Path,
        default=REPOSITORY_ROOT / "output/3s-submission-evidence",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=REPOSITORY_ROOT / "output/3s-submission-documents",
    )
    parser.add_argument("--approvals", type=Path)
    parser.add_argument("--pdftoppm", type=Path)
    parser.add_argument("--pdfinfo", type=Path)
    parser.add_argument("--word-timeout-seconds", type=float, default=60)
    parser.add_argument("--verify-only", action="store_true")
    args = parser.parse_args(argv)
    try:
        load_verified_submission_metadata(args.metadata, REPOSITORY_ROOT)
        if args.verify_only:
            if not _verify_existing_documents(args.output_dir, args.approvals):
                raise DocumentGenerationError("文档或逐页哈希视觉批准门未通过。")
        else:
            pdftoppm = args.pdftoppm or (Path(value) if (value := shutil.which("pdftoppm")) else None)
            pdfinfo = args.pdfinfo or (Path(value) if (value := shutil.which("pdfinfo")) else None)
            if pdftoppm is None or pdfinfo is None:
                raise DocumentGenerationError("必须显式提供 Poppler 路径或将其加入 PATH。")
            generate_documents(
                metadata_path=args.metadata,
                evidence_directory=args.evidence_dir,
                output_directory=args.output_dir,
                pdftoppm=pdftoppm,
                pdfinfo=pdfinfo,
                approvals_path=args.approvals,
                word_timeout_seconds=args.word_timeout_seconds,
            )
    except (DocumentGenerationError, SubmissionGateError, OSError) as exc:
        print(str(exc), file=sys.stderr)
        return 1
    print("3S 申报文档已生成并完成结构与渲染检查。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
